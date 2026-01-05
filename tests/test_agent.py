from pathlib import Path
from typing import List

import pytest
from fpdf import FPDF
from langchain_community.embeddings import FakeEmbeddings
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from rag_agent.agent import VisaAgent, build_retriever, load_documents
from docx import Document


class StubLLM(BaseChatModel):
    @property
    def _llm_type(self) -> str:
        return "stub"

    @property
    def _identifying_params(self) -> dict:
        return {}

    def _generate(
        self,
        messages: List[HumanMessage],
        stop: List[str] | None = None,
        run_name: str | None = None,
        **kwargs,
    ) -> ChatResult:
        last_human = next(
            (message.content for message in reversed(messages) if isinstance(message, HumanMessage)),
            "",
        )
        return ChatResult(
            generations=[ChatGeneration(message=AIMessage(content=f"Stub answer using: {last_human}"))]
        )


def _create_pdf(path: Path, text: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.splitlines():
        pdf.cell(200, 10, txt=line, ln=1)
    pdf.output(str(path))


def _create_docx(path: Path, text: str) -> None:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def test_load_documents_reads_pdf_and_docx(tmp_path: Path) -> None:
    pdf_path = tmp_path / "guia.pdf"
    docx_path = tmp_path / "roteiro.docx"
    _create_pdf(pdf_path, "Requisitos do visto D7: comprovativo de rendimentos.")
    _create_docx(docx_path, "Benefícios do visto D7 incluem residência e reagrupamento.")

    docs = load_documents(tmp_path)

    assert len(docs) == 2
    assert any("Requisitos" in doc.page_content for doc in docs)
    assert any("Benefícios" in doc.page_content for doc in docs)


def test_agent_uses_context_and_stubbed_llm(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    _create_pdf(data_dir / "guia.pdf", "Requisitos mínimos para o Visto D7: renda passiva.")
    _create_docx(data_dir / "roteiro.docx", "Prazos e etapas para obter o visto D7.")

    fake_embeddings = FakeEmbeddings(size=32)
    retriever = build_retriever(
        data_dir=data_dir,
        embedding_model=fake_embeddings,
        persist_dir=tmp_path / "persist",
        reindex=True,
    )
    agent = VisaAgent(
        data_dir=data_dir,
        persist_dir=tmp_path / "persist",
        reindex=False,
        llm=StubLLM(),
        embedding=fake_embeddings,
        search_client=None,
    )
    agent.retriever = retriever  # ensure we reuse the fake retriever

    response = agent.answer_question("Quais são os requisitos do visto?")

    assert "Visto D7" in response
    assert "Requisitos mínimos" in response


def test_build_retriever_returns_none_when_no_documents(tmp_path: Path) -> None:
    retriever = build_retriever(
        data_dir=tmp_path,
        embedding_model=FakeEmbeddings(size=16),
        persist_dir=tmp_path / "persist",
        reindex=True,
    )

    assert retriever is None
